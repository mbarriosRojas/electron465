#!/usr/bin/env python3
"""Clientes con credito 100% cancelado y sin contratos activos, clasificados ORO / PLATA / BRONCE.
Uso: generar_reporte.py [--db copy_electron] [--out DIR]"""
import argparse, collections as co, csv, datetime as dt, io, os, subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ap = argparse.ArgumentParser()
ap.add_argument('--db', default='copy_electron')
ap.add_argument('--out', default='/Volumes/goku/dev_php/db_backups')
a = ap.parse_args()

def q(sql):
    r = subprocess.run(['mysql', '-uroot', '--default-character-set=utf8mb4', a.db, '-B', '-e', sql], capture_output=True, text=True, check=True)
    return list(csv.DictReader(io.StringIO(r.stdout), delimiter='\t', quoting=csv.QUOTE_NONE))

TOL = 0.01
# Umbrales de nivel
ORO = dict(extra=1, maxat=1.5, pct=0.85)
BRONCE = dict(extra=6, maxat=4, pct=0.5)
AL_DIA = 1.05  # atraso <= 1 cuota cuenta como al dia (el banco cobra con desfase)

C = q("""SELECT c.contrato_id, c.documento_id, c.fecha_inicio_cobro, c.numero_cuotas, c.monto_cuota, c.monto_total,
  c.periocidad, c.cobrado_en, c.nomina_procedencia,
  TRIM(CONCAT_WS(' ', NULLIF(TRIM(p.primer_nombre),''), NULLIF(TRIM(p.segundo_nombre),''), NULLIF(TRIM(p.primer_apellido),''), NULLIF(TRIM(p.segundo_apellido),''))) nombre,
  IFNULL(p.banco_1,'') banco_1
  FROM t_clientes_creditos c LEFT JOIN t_estadoejecucion e ON e.oidc=c.contrato_id
  LEFT JOIN t_personas p ON p.documento_id=c.documento_id
  WHERE c.cantidad<>0 AND c.estatus<>3 AND (e.oide IS NULL OR e.oide<>0)""")
pays = co.defaultdict(lambda: co.defaultdict(float))
for p in q("SELECT credito_id, CAST(fecha AS CHAR) fecha, monto, mesp, anop FROM t_lista_cobros"):
    try: y, mo = int(p['anop']), int(p['mesp'])
    except ValueError: y = mo = 0
    if not (2000 < y < 2100 and 1 <= mo <= 12):
        f = p['fecha']
        if not f or f.startswith('0000') or f == 'NULL': continue
        y, mo = int(f[:4]), int(f[5:7])
    pays[p['credito_id']][y * 12 + mo - 1] += float(p['monto'] or 0)

CPM = {'0': 52 / 12, '1': 26 / 12, '2': 2, '3': 2}
PERIO = {'0': 'SEMANAL', '1': 'CATORCENAL', '2': 'QUINCENAL 15-30', '3': 'QUINCENAL 10-20', '4': 'MENSUAL', '8': 'MENSUAL-10', '9': 'MENSUAL-25'}
RANK = {'ORO': 0, 'PLATA': 1, 'BRONCE': 2}
ym = lambda m: f"{m // 12}-{m % 12 + 1:02d}"

def analizar(c):
    bym = pays[c['contrato_id']]; tot = float(c['monto_total']); cu = float(c['monto_cuota'])
    f = c['fecha_inicio_cobro']; m0 = int(f[:4]) * 12 + int(f[5:7]) - 1
    cpm = CPM.get(c['periocidad'], 1); permes = cu * cpm
    mend = m0 + max(0, int(-(-int(c['numero_cuotas']) // cpm)) - 1)
    first = min([m0] + list(bym)); last = max(bym) if bym else m0
    cum = 0; payoff = None; maxat = 0; ald = nexp = sinpago = picos = 0
    for mm in range(first, max(last, mend) + 1):
        pm = bym.get(mm, 0); cum += pm
        if payoff is None and cum >= tot - TOL: payoff = mm
        if mm < m0: continue
        if payoff is not None and mm > payoff: break
        at = max(0, (min(tot, (mm - m0 + 1) * permes) - cum) / cu) if cu > 0 else 0
        nexp += 1; ald += at <= AL_DIA; maxat = max(maxat, at)
        sinpago += pm <= 0.01; picos += pm > 1.5 * permes
    payoff = payoff if payoff is not None else last
    r = dict(extra=payoff - mend, maxat=maxat, pct=ald / nexp if nexp else 1, sinpago=sinpago, picos=picos, mend=mend, payoff=payoff)
    if r['extra'] <= ORO['extra'] and r['maxat'] <= ORO['maxat'] and r['pct'] >= ORO['pct']: r['nivel'] = 'ORO'
    elif r['extra'] > BRONCE['extra'] or r['maxat'] > BRONCE['maxat'] or r['pct'] < BRONCE['pct']: r['nivel'] = 'BRONCE'
    else: r['nivel'] = 'PLATA'
    return r

byc = co.defaultdict(list)
for c in C:
    c['paid'] = sum(pays[c['contrato_id']].values()); c['saldo'] = float(c['monto_total']) - c['paid']
    byc[c['documento_id']].append(c)

ced = lambda d: int(d) if d.isdigit() else d
cnt = co.Counter(); clientes = []; detalle = []
for d, cs in byc.items():
    if any(c['saldo'] > TOL for c in cs): continue
    cs.sort(key=lambda c: (c['fecha_inicio_cobro'], c['contrato_id'])); rr = [analizar(c) for c in cs]
    nivel = max((r['nivel'] for r in rr), key=RANK.get); cnt[nivel] += 1; last = cs[-1]
    clientes.append([ced(d), last['nombre'], last['nomina_procedencia'], last['cobrado_en'], last['banco_1'], nivel, len(cs),
        round(sum(c['paid'] for c in cs), 2), last['contrato_id'], cs[0]['fecha_inicio_cobro'], ym(rr[-1]['payoff']),
        max(0, max(r['extra'] for r in rr)), round(max(r['maxat'] for r in rr), 1), round(sum(r['pct'] for r in rr) / len(rr), 3),
        sum(r['sinpago'] for r in rr), sum(r['picos'] for r in rr)])
    for c, r in zip(cs, rr):
        detalle.append([ced(d), c['nombre'], c['contrato_id'], r['nivel'], c['nomina_procedencia'], c['cobrado_en'], PERIO.get(c['periocidad'], c['periocidad']),
            int(c['numero_cuotas']), float(c['monto_cuota']), float(c['monto_total']), round(c['paid'], 2), round(c['saldo'], 2), c['fecha_inicio_cobro'],
            ym(r['mend']), ym(r['payoff']), max(0, r['extra']), round(r['maxat'], 1), round(r['pct'], 3), r['sinpago'], r['picos']])
clientes.sort(key=lambda x: (RANK[x[5]], x[1]))

FILL = {'ORO': 'FFD966', 'PLATA': 'D9D9D9', 'BRONCE': 'F4B183'}
def hoja(ws, hdr, rows, lv, pct):
    ws.append(hdr)
    for r in rows: ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78'); c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        row[lv].fill = PatternFill('solid', fgColor=FILL[row[lv].value]); row[lv].font = Font(bold=True); row[pct].number_format = '0.0%'
    for i, col in enumerate(ws.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, max(len(str(c.value or '')) for c in list(col)[:300]) + 2), 45)

wb = Workbook(); rs = wb.active; rs.title = 'Resumen'
hoja(wb.create_sheet('Clientes'), ['Cédula', 'Nombre', 'Nómina', 'Banco (cobrado en)', 'Banco cuenta cliente', 'Nivel', 'Contratos cancelados',
    'Monto total pagado (USD)', 'Último contrato', 'Fecha inicio 1er contrato', 'Mes cancelación último', 'Meses de retraso (peor)',
    'Máx. cuotas atrasadas', '% meses al día (prom.)', 'Meses sin pago (total)', 'Pagos en pico (total)'], clientes, 5, 13)
hoja(wb.create_sheet('Detalle contratos'), ['Cédula', 'Nombre', 'Contrato', 'Nivel contrato', 'Nómina', 'Banco (cobrado en)', 'Periodicidad', 'N° cuotas',
    'Monto cuota', 'Monto total', 'Pagado', 'Saldo', 'Inicio cobro', 'Fin esperado', 'Mes cancelación', 'Meses de retraso', 'Máx. cuotas atrasadas',
    '% meses al día', 'Meses sin pago', 'Pagos en pico'], detalle, 3, 17)
now = dt.datetime.now()
for l in [['Clientes con crédito 100% cancelado y sin contratos activos'], [f'Base: {a.db} — generado {now:%Y-%m-%d %H:%M}'], [], ['Nivel', 'Clientes'],
    *[[k, cnt[k]] for k in RANK], ['TOTAL', sum(cnt.values())], [], ['Criterios'],
    ['Cliente elegible', 'Todos sus contratos válidos (no nulos ni rechazados en ventas) tienen saldo 0: monto_total − cobros ≤ 0,01 USD. Si tiene cualquier contrato con saldo pendiente, no aparece.'],
    ['Nivel del cliente', 'Se toma el PEOR nivel entre sus contratos cancelados (detalle en la hoja "Detalle contratos").'],
    ['ORO', 'Terminó como máximo 1 mes después de lo pactado, nunca acumuló más de 1,5 cuotas de atraso y estuvo al día (≤1 cuota de atraso) en al menos 85% de los meses.'],
    ['PLATA', 'No cumple ORO ni BRONCE: se extendió hasta 6 meses o tuvo atrasos moderados.'],
    ['BRONCE', 'Se extendió más de 6 meses, o llegó a deber más de 4 cuotas, o estuvo al día en menos de la mitad de los meses (pagos irregulares / en picos).'],
    [], ['Glosario'],
    ['Meses de retraso', 'Meses entre el mes de la última cuota pactada y el mes en que se completó el pago.'],
    ['Máx. cuotas atrasadas', 'Mayor diferencia, en cuotas, entre lo que debía llevar pagado y lo pagado en algún mes.'],
    ['% meses al día', 'Proporción de meses (desde el inicio de cobro hasta la cancelación) con atraso ≤ 1 cuota.'],
    ['Meses sin pago', 'Meses del período de cobro en que no ingresó ningún monto.'],
    ['Pagos en pico', 'Meses en que pagó más de 1,5 veces la cuota mensual (pagos de recuperación).']]:
    rs.append(l)
rs['A1'].font = Font(bold=True, size=14)
for r in (4, 9, 16): rs.cell(r, 1).font = Font(bold=True)
for i, k in enumerate(RANK, 5): rs.cell(i, 1).fill = PatternFill('solid', fgColor=FILL[k])
rs.column_dimensions['A'].width = 24; rs.column_dimensions['B'].width = 110
for row in rs.iter_rows(min_row=10): row[1].alignment = Alignment(wrap_text=True, vertical='top')

out = os.path.join(a.out, f'reporte_clientes_cancelados_nivel_{now:%Y%m%d_%H%M}.xlsx')
wb.save(out)
print(out); print(dict(ORO=cnt['ORO'], PLATA=cnt['PLATA'], BRONCE=cnt['BRONCE'], TOTAL=sum(cnt.values())))
print('datos hasta:', q("SELECT MAX(modificado) m FROM t_lista_cobros")[0]['m'])
